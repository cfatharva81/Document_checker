"""build_doc() is the extractor's entry point.

This file just wires the other pieces together: open the .docx, resolve its
styles, walk the body, infer headings if Word didn't record any, read the
sections, and read the document's core properties (title, author, etc).
"""
from __future__ import annotations

from typing import Optional

from docx import Document as OpenDocument
from lxml import etree

from .headings import infer_heading_levels
from .model import CoreProps, Doc
from .sections import build_sections
from .styles import StyleResolver
from .walker import Ctx, walk_container

THEME_RELTYPE = ("http://schemas.openxmlformats.org/officeDocument/2006/"
                 "relationships/theme")


def build_doc(file, filename: str) -> Doc:
    document = OpenDocument(file)
    resolver = StyleResolver(_styles_element(document), _theme_element(document))

    ctx = Ctx(resolver)
    body = document.element.body
    blocks = walk_container(body, ctx)

    infer_heading_levels(blocks, ctx.tables)

    sections = build_sections(document, resolver)
    core = _core_props(document)

    return Doc(
        filename=filename,
        core=core,
        blocks=blocks,
        tables=ctx.tables,
        sections=sections,
        styles=resolver.index(),
    )


# --------------------------------------------------------------------------
# Part helpers
# --------------------------------------------------------------------------
def _styles_element(document) -> Optional[etree._Element]:
    try:
        return document.styles.element
    except Exception:
        return None


def _theme_element(document) -> Optional[etree._Element]:
    try:
        for rel in document.part.rels.values():
            if rel.reltype == THEME_RELTYPE:
                try:
                    return etree.fromstring(rel.target_part.blob)
                except Exception:
                    return None
    except Exception:
        return None
    return None


def _core_props(document) -> CoreProps:
    cp = document.core_properties
    return CoreProps(
        title=cp.title or None,
        author=cp.author or None,
        subject=cp.subject or None,
        keywords=cp.keywords or None,
        created=cp.created,
        modified=cp.modified,
        last_modified_by=cp.last_modified_by or None,
    )
